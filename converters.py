import subprocess
import csv
import json
import shutil
import os
from tempfile import TemporaryDirectory
from pathlib import Path
from docx import Document

from core import converter

def run_libreoffice(source: Path, work_dir: Path, target_ext: str,
                    executable: str | None = None) -> Path:
    tool = (executable or shutil.which("libreoffice")
            or shutil.which("soffice"))
    if not tool:
        windows_path = r"C:\Program Files\LibreOffice\program\soffice.exe"
        if os.path.exists(windows_path):
            tool = windows_path
        else:
            raise RuntimeError("Brak LibreOffice. Zainstaluj program lub ustaw ścieżkę.")

    command = [tool, "--headless", "--convert-to", target_ext,
               "--outdir", str(work_dir), str(source.resolve())]

    try:
        result = subprocess.run(
            command, shell=False, capture_output=True,
            text=True, errors="replace", timeout=60
        )
    except FileNotFoundError as exc:
        raise RuntimeError("Sprawdź ścieżkę do LibreOffice") from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("Przekroczono czas konwersji (60 s)") from exc
    except OSError as exc:
        raise RuntimeError("Nie można uruchomić LibreOffice") from exc

    if result.returncode != 0:
        detail = result.stderr.strip() or result.stdout.strip()
        raise RuntimeError(detail or "Błąd procesu LibreOffice")

    target = work_dir / f"{source.stem}.{target_ext}"
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError(f"Nie powstał niepusty plik {target_ext.upper()}")
    return target


@converter(".docx", ".pdf")
def docx_to_pdf(source: Path, output_dir: Path,
                executable: str | None = None) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if source.suffix.lower() != ".docx":
        raise ValueError("Oczekiwano pliku DOCX")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")

    target = output_dir / (source.stem + ".pdf")
    if target.exists():
        raise FileExistsError("Plik PDF już istnieje")

    with TemporaryDirectory() as temp:
        result_in_temp = run_libreoffice(source, Path(temp), "pdf", executable)
        with result_in_temp.open("rb") as src, target.open("xb") as dst:
            shutil.copyfileobj(src, dst)

    return target


@converter(".txt", ".docx")
def txt_to_docx(source: Path, output_dir: Path) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")
    if source.suffix.lower() != ".txt":
        raise ValueError("Oczekiwano pliku TXT")

    target = output_dir / (source.stem + ".docx")
    if target.exists():
        raise FileExistsError("Wynik istnieje; wybierz inną nazwę")

    text = source.read_text(encoding="utf-8")
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    document.save(str(target))
    return target


@converter(".docx", ".txt")
def docx_to_txt(source: Path, output_dir: Path) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")
    if source.suffix.lower() != ".docx":
        raise ValueError("Oczekiwano pliku DOCX")

    target = output_dir / (source.stem + ".txt")
    if target.exists():
        raise FileExistsError("Wynik istnieje; wybierz inną nazwę")

    doc = Document(source)
    text = "\n".join(para.text for para in doc.paragraphs)

    target.write_text(text, encoding="utf-8")
    return target


@converter(".csv", ".json")
def csv_to_json(source: Path, output_dir: Path) -> Path:
    target: Path = output_dir / f"{source.stem}.json"

    output_dir.mkdir(parents=True, exist_ok=True)

    with open(source, mode='r', encoding="utf-8") as csv_file:
        reader = csv.DictReader(csv_file)
        data = list(reader)

    with open(target, mode="w", encoding="utf-8") as json_file:
        json.dump(data, json_file, ensure_ascii=False, indent=4)

    return target


@converter(".json", ".csv")
def json_to_csv(source: Path, output_dir: Path) -> Path:
    target = output_dir / f"{source.stem}.csv"

    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        with open(source, mode="r", encoding="utf-8") as json_file:
            data = json.load(json_file)
    except json.JSONDecodeError:
        raise ValueError(f"Błąd: plik {source.name} nie jest poprawnym plikiem JSON. Plik jest uszkodzony lub błędnie sformatowany.")

    if len(data) == 0:
        raise ValueError("Plik JSON jest pusty, brak danych do konwersji.")

    first_row: list[str] = list(data[0].keys())

    with open(target, mode="w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=first_row)
        writer.writeheader()
        writer.writerows(data)

    return target


@converter(".docx", ".odt")
def docx_to_odt(source: Path, output_dir: Path,
                executable: str | None = None) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if source.suffix.lower() != ".docx":
        raise ValueError("Oczekiwano pliku DOCX")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")

    target = output_dir / (source.stem + ".odt")
    if target.exists():
        raise FileExistsError("Plik ODT już istnieje")

    with TemporaryDirectory() as temp:
        result_in_temp = run_libreoffice(source, Path(temp), "odt", executable)
        with result_in_temp.open("rb") as src, target.open("xb") as dst:
            shutil.copyfileobj(src, dst)

    return target


@converter(".odt", ".docx")
def odt_to_docx(source: Path, output_dir: Path,
                executable: str | None = None) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if source.suffix.lower() != ".odt":
        raise ValueError("Oczekiwano pliku ODT")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")

    target = output_dir / (source.stem + ".docx")
    if target.exists():
        raise FileExistsError("Plik DOCX już istnieje")

    with TemporaryDirectory() as temp:
        result_in_temp = run_libreoffice(source, Path(temp), "docx", executable)
        with result_in_temp.open("rb") as src, target.open("xb") as dst:
            shutil.copyfileobj(src, dst)

    return target