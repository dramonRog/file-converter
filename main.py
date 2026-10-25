from dataclasses import dataclass
from pathlib import Path
from collections.abc import Callable
import subprocess
import csv
import json
import shutil
from tempfile import TemporaryDirectory
import os
from docx import Document
import time


Converter = Callable[[Path, Path], Path]
REGISTRY: dict[tuple[str, str], Converter] = {}


def converter(source: str, target: str):
    """Dekorator rejestrujący pojedynczy kierunek konwersji."""
    def register(func: Converter) -> Converter:
        REGISTRY[(source.lower(), target.lower())] = func
        return func
    return register


@dataclass(frozen=True)
class ConversionRequest:
    source: Path
    target_format: str
    output_dir: Path

    def validate(self) -> None:
        if not self.source.is_file():
            raise FileNotFoundError("Nie znaleziono pliku wejściowego")
        if not self.output_dir.is_dir():
            raise NotADirectoryError("Nie znaleziono katalogu docelowego")
        key = (self.source.suffix.lower(), self.target_format.lower())
        if key not in REGISTRY:
            raise ValueError(f"Nieobsługiwana konwersja: {key}")


def available_targets(source: Path) -> list[str]:
    return sorted(target for (src, target) in REGISTRY if src == source.suffix.lower())


def run_pdf(source: Path, work_dir: Path,
            executable: str | None = None) -> Path:
    tool = (executable or shutil.which("libreoffice")
            or shutil.which("soffice"))
    if not tool:
        windows_path = r"C:\Program Files\LibreOffice\program\soffice.exe"
        if os.path.exists(windows_path):
            tool = windows_path
        else:
            raise RuntimeError("Brak LibreOffice. Zainstaluj program lub ustaw ścieżkę.")

    command = [tool, "--headless", "--convert-to", "pdf",
               "--outdir", str(work_dir), str(source.resolve())]

    try:
        result = subprocess.run(
            command, shell=False, capture_output=True,
            text=True, errors="replace", timeout=60
        )
    except FileNotFoundError as exc:
        raise RuntimeError("Sprawdź ścieżkę do LibreOffice") from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("Przekroczono czas konwersji (60 s)") from exc
    except OSError as exc:
        raise RuntimeError("Nie można uruchomić LibreOffice") from exc

    if result.returncode != 0:
        detail = result.stderr.strip() or result.stdout.strip()
        raise RuntimeError(detail or "Błąd procesu LibreOffice")

    target = work_dir / (source.stem + ".pdf")
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError("Nie powstał niepusty plik PDF")
    return target


@converter(".docx", ".pdf")
def docx_to_pdf(source: Path, output_dir: Path,
                executable: str | None = None) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if source.suffix.lower() != ".docx":
        raise ValueError("Oczekiwano pliku DOCX")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")

    target = output_dir / (source.stem + ".pdf")
    if target.exists():
        raise FileExistsError("Plik PDF już istnieje")

    with TemporaryDirectory() as temp:
        result_in_temp = run_libreoffice(source, Path(temp), "pdf", executable)
        with result_in_temp.open("rb") as src, target.open("xb") as dst:
            shutil.copyfileobj(src, dst)

    return target


@converter(".txt", ".docx")
def txt_to_docx(source: Path, output_dir: Path) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")
    if source.suffix.lower() != ".txt":
        raise ValueError("Oczekiwano pliku TXT")

    target = output_dir / (source.stem + ".docx")
    if target.exists():
        raise FileExistsError("Wynik istnieje; wybierz inną nazwę")

    text = source.read_text(encoding="utf-8")
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    document.save(str(target))
    return target


@converter(".docx", ".txt")
def docx_to_txt(source: Path, output_dir: Path) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")
    if source.suffix.lower() != ".docx":
        raise ValueError("Oczekiwano pliku DOCX")

    target = output_dir / (source.stem + ".txt")
    if target.exists():
        raise FileExistsError("Wynik istnieje; wybierz inną nazwę")

    doc = Document(source)
    text = "\n".join(para.text for para in doc.paragraphs)

    target.write_text(text, encoding="utf-8")
    return target


@converter(".csv", ".json")
def csv_to_json(source: Path, output_dir: Path) -> Path:
    target: Path = output_dir / f"{source.stem}.json"

    output_dir.mkdir(parents=True, exist_ok=True)

    with open(source, mode='r', encoding="utf-8") as csv_file:
        reader = csv.DictReader(csv_file)
        data = list(reader)

    with open(target, mode="w", encoding="utf-8") as json_file:
        json.dump(data, json_file, ensure_ascii=False, indent=4)

    return target


@converter(".json", ".csv")
def json_to_csv(source: Path, output_dir: Path) -> Path:
    target = output_dir / f"{source.stem}.csv"

    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        with open(source, mode="r", encoding="utf-8") as json_file:
            data = json.load(json_file)
    except json.JSONDecodeError:
        raise ValueError(f"Błąd: plik {source.name} nie jest poprawnym plikiem JSON. Plik jest uszkodzony lub błędnie sformatowany.")

    if len(data) == 0:
        raise ValueError("Plik JSON jest pusty, brak danych do konwersji.")

    first_row: list[str] = list(data[0].keys())

    with open(target, mode="w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=first_row)
        writer.writeheader()
        writer.writerows(data)

    return target


@converter(".docx", ".odt")
def docx_to_odt(source: Path, output_dir: Path,
                executable: str | None = None) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if source.suffix.lower() != ".docx":
        raise ValueError("Oczekiwano pliku DOCX")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")

    target = output_dir / (source.stem + ".odt")
    if target.exists():
        raise FileExistsError("Plik ODT już istnieje")

    with TemporaryDirectory() as temp:
        result_in_temp = run_libreoffice(source, Path(temp), "odt", executable)
        with result_in_temp.open("rb") as src, target.open("xb") as dst:
            shutil.copyfileobj(src, dst)

    return target


@converter(".odt", ".docx")
def odt_to_docx(source: Path, output_dir: Path,
                executable: str | None = None) -> Path:
    if not source.is_file():
        raise FileNotFoundError("Brak pliku wejściowego")
    if source.suffix.lower() != ".odt":
        raise ValueError("Oczekiwano pliku ODT")
    if not output_dir.is_dir():
        raise NotADirectoryError("Brak katalogu wynikowego")

    target = output_dir / (source.stem + ".docx")
    if target.exists():
        raise FileExistsError("Plik DOCX już istnieje")

    with TemporaryDirectory() as temp:
        result_in_temp = run_libreoffice(source, Path(temp), "docx", executable)
        with result_in_temp.open("rb") as src, target.open("xb") as dst:
            shutil.copyfileobj(src, dst)

    return target


def run_libreoffice(source: Path, work_dir: Path, target_ext: str,
                    executable: str | None = None) -> Path:
    tool = (executable or shutil.which("libreoffice")
            or shutil.which("soffice"))
    if not tool:
        windows_path = r"C:\Program Files\LibreOffice\program\soffice.exe"
        if os.path.exists(windows_path):
            tool = windows_path
        else:
            raise RuntimeError("Brak LibreOffice. Zainstaluj program lub ustaw ścieżkę.")

    command = [tool, "--headless", "--convert-to", target_ext,
               "--outdir", str(work_dir), str(source.resolve())]

    try:
        result = subprocess.run(
            command, shell=False, capture_output=True,
            text=True, errors="replace", timeout=60
        )
    except FileNotFoundError as exc:
        raise RuntimeError("Sprawdź ścieżkę do LibreOffice") from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("Przekroczono czas konwersji (60 s)") from exc
    except OSError as exc:
        raise RuntimeError("Nie można uruchomić LibreOffice") from exc

    if result.returncode != 0:
        detail = result.stderr.strip() or result.stdout.strip()
        raise RuntimeError(detail or "Błąd procesu LibreOffice")

    target = work_dir / f"{source.stem}.{target_ext}"
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError(f"Nie powstał niepusty plik {target_ext.upper()}")
    return target


def execute(request: ConversionRequest) -> Path:
    request.validate()
    key = (request.source.suffix.lower(), request.target_format.lower())

    start_time = time.perf_counter()
    try:
        result_path = REGISTRY[key](request.source, request.output_dir)

        elapsed_time = time.perf_counter() - start_time
        print(f"Konwersja {request.source.name} -> {request.target_format} trwała {elapsed_time:.2f} s")

        return result_path
    except Exception as e:
        print(f"Błąd podczas konwersji {request.source.name}: {str(e)}")
        raise


if __name__ == "__main__":
    test_source_csv = Path("test/csv/test_csv.csv")
    test_source_json = Path("test/json/test_json.json")
    test_error_json = Path("test/json/error-json.json")
    test_source_txt = Path("test/txt/test_text.txt")
    test_source_docx = Path("test/docx/test_text1.docx")
    test_source_odt = Path("test/odt/test_odt.odt")

    output_dir_json = Path("test/json")
    output_dir_csv = Path("test/csv")
    output_dir_pdf = Path("test/pdf")
    output_dir_docx = Path("test/docx")
    output_dir_txt = Path("test/txt")
    output_dir_odt = Path("test/odt")

    if test_source_csv.exists():
        req_csv_to_json = ConversionRequest(
            source=test_source_csv,
            target_format=".json",
            output_dir=output_dir_json
        )
        json_result = execute(req_csv_to_json)
        print(f"CSV do JSON: zapisano w {json_result}")
    else:
        print(f"Utwórz plik {test_source_csv}. Plik nie był znaleziony.")

    if test_error_json.exists():
        req_json_to_csv = ConversionRequest(
            source=test_error_json,
            target_format=".csv",
            output_dir=output_dir_csv
        )
        try:
            csv_result = execute(req_json_to_csv)
        except ValueError as ex:
            print(ex)

    if test_source_json.exists():
        req_json_to_csv = ConversionRequest(
            source=test_source_json,
            target_format=".csv",
            output_dir=output_dir_csv
        )
        csv_result = execute(req_json_to_csv)
        print(f"JSON to CSV: zapisano w {csv_result}")
    else:
        print(f"Utwórz plik {test_source_json}. Plik nie był znaleziony.")

    test_source_docx = Path("test/docx/mini-test-doc.docx")

    if test_source_docx.exists():
        req_docx_to_pdf = ConversionRequest(
            source=test_source_docx,
            target_format=".pdf",
            output_dir=output_dir_pdf
        )
        try:
            pdf_result = execute(req_docx_to_pdf)
            print(f"DOCX do PDF: zapisano w {pdf_result}")
        except Exception as e:
            print(f"Błąd konwersji DOCX do PDF: {e}")
    else:
        print(f"Utwórz plik {test_source_docx}. Plik nie był znaleziony.")

    if test_source_txt.exists():
        req_txt_to_docx = ConversionRequest(
            source=test_source_txt,
            target_format=".docx",
            output_dir=output_dir_docx
        )
        try:
            docx_result = execute(req_txt_to_docx)
            print(f"TXT do DOCX: zapisano w {docx_result}")
        except Exception as e:
            print(f"Błąd konwersji TXT do DOCX: {e}")
    else:
        print(f"Utwórz plik {test_source_txt}. Plik nie był znaleziony.")

    if test_source_docx.exists():
        req_docx_to_txt = ConversionRequest(
            source=test_source_docx,
            target_format=".txt",
            output_dir=output_dir_txt
        )
        try:
            txt_result = execute(req_docx_to_txt)
            print(f"DOCX do TXT: zapisano w {txt_result}")
        except Exception as e:
            print(f"Błąd konwersji DOCX do TXT: {e}")

    if test_source_docx.exists():
        req_docx_to_odt = ConversionRequest(
            source=test_source_docx,
            target_format=".odt",
            output_dir=output_dir_odt
        )
        try:
            odt_result = execute(req_docx_to_odt)
            print(f"DOCX do ODT: zapisano w {odt_result}")
        except Exception as e:
            print(f"Błąd konwersji DOCX do ODT: {e}")

    if test_source_odt.exists():
        req_odt_to_docx = ConversionRequest(
            source=test_source_odt,
            target_format=".docx",
            output_dir=output_dir_docx
        )
        try:
            docx_result_from_odt = execute(req_odt_to_docx)
            print(f"ODT do DOCX: zapisano w {docx_result_from_odt}")
        except Exception as e:
            print(f"Błąd konwersji ODT do DOCX: {e}")
    else:
        print(f"Utwórz plik {test_source_odt}. Plik nie był znaleziony.")